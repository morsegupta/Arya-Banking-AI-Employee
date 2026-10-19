"""Spreadsheet tools and the tracker sync: only writable columns, never over an open workbook,
people's columns preserved, every write read back."""

from __future__ import annotations

import asyncio

from openpyxl import load_workbook

from apps.tracker import sync_trackers
from sandbox import chaos
from tests.fakes import ScriptedLLM
from tools.base import RunContext


def _row(sandbox_var, ref):
    ws = load_workbook(sandbox_var / "drive" / "finance" / "ar_tracker.xlsx")["Receivables"]
    header = [c.value for c in ws[1]]
    return next(dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if r[0].value == ref)


def _invoke(rt, tool, args):
    ctx = RunContext(mission_id="mis_t", task_id="tsk_t", allowed_apps={"ar_tracker"}, artifacts_dir=rt.artifacts_dir)
    return asyncio.run(rt.registry.invoke(tool, args, ctx))


def test_update_writes_only_writable_columns_and_reads_back(make_runtime, sandbox_var):
    rt = make_runtime(ScriptedLLM([]))
    res = _invoke(rt, "ar_tracker.update_row", {"key": "INV-2291", "values": {"Status": "Promised", "Balance (₹)": 340000}})
    assert res.ok and res.verified is True, res.error
    row = _row(sandbox_var, "INV-2291")
    assert row["Status"] == "Promised" and row["Notes"] == "Ritu says payment run is on Fridays"
    denied = _invoke(rt, "ar_tracker.update_row", {"key": "INV-2291", "values": {"Notes": "overwritten"}})
    assert not denied.ok and denied.error_kind == "denied"
    missing = _invoke(rt, "ar_tracker.update_row", {"key": "INV-9999", "values": {"Status": "Paid"}})
    assert not missing.ok and missing.error_kind == "permanent"
    asyncio.run(rt.close())


def test_an_open_workbook_is_never_overwritten(make_runtime, sandbox_var):
    rt = make_runtime(ScriptedLLM([]))
    chaos.set_switches(**{"tracker.open_in_app": True})
    res = _invoke(rt, "ar_tracker.update_row", {"key": "INV-2291", "values": {"Status": "Paid"}})
    assert not res.ok and res.error_kind == "transient" and "open in the desktop app" in res.error
    assert _row(sandbox_var, "INV-2291")["Status"] == "Open"
    chaos.set_switches(**{"tracker.open_in_app": False})
    assert _invoke(rt, "ar_tracker.update_row", {"key": "INV-2291", "values": {"Status": "Paid"}}).ok
    asyncio.run(rt.close())


def test_sync_follows_the_pipeline_and_only_changed_rows(make_runtime, sandbox_var):
    rt = make_runtime(ScriptedLLM([]))
    e = rt.pipeline.add_entity("mis_t", "INV-2297", state="OVERDUE", value=5_400_000,
                               data={"balance": 5_400_000, "record": {"customer_id": "C11"}})
    first = asyncio.run(sync_trackers(rt, "mis_t"))
    assert first["updated"] == ["INV-2297"] and _row(sandbox_var, "INV-2297")["Status"] == "Overdue"
    assert asyncio.run(sync_trackers(rt, "mis_t"))["updated"] == [], "nothing changed: nothing written"
    rt.pipeline.propose(e.id, "CONTACTED", {"outreach": {"delivered": True, "message_id": "msg_1"}})
    rt.pipeline.update_data(e.id, {"last_contact": {"channel": "whatsapp"}})
    assert asyncio.run(sync_trackers(rt, "mis_t"))["updated"] == ["INV-2297"]
    row = _row(sandbox_var, "INV-2297")
    assert row["Status"] == "Reminder sent" and row["Last action"] == "Reminder sent by whatsapp"
    assert row["Updated"].endswith("by Arya") and row["Balance (₹)"] == 54000
    assert any(x["subject"] == "tracker:ar_tracker" for x in rt.evidence.for_entity(e.id))
    asyncio.run(rt.close())
