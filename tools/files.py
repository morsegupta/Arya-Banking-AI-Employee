"""File tools: read documents from the shared drive and from the task's own files.

Paths are confined to two roots: the company's shared drive and the current
task's artifact folder (downloads, attachments). `..` tricks and symlinks
that escape a root are refused. Document contents are third-party data and
are marked untrusted (a PDF can carry hidden instructions).
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import ClassVar

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from core.models import ToolResult
from tools.base import PermanentToolError, RunContext, Tool, ToolDenied, ToolRegistry


class FileRoots:
    def __init__(self, drive_root: Path | str):
        self.drive_root = Path(drive_root).resolve()

    def roots(self, ctx: RunContext) -> list[Path]:
        return [self.drive_root, ctx.task_dir().resolve()]

    def resolve(self, raw: str, ctx: RunContext) -> Path:
        """Relative paths are relative to the drive. Result must stay inside an allowed root."""
        candidate = Path(raw)
        path = (candidate if candidate.is_absolute() else self.drive_root / candidate).resolve()
        if not any(path == r or path.is_relative_to(r) for r in self.roots(ctx)):
            raise ToolDenied(f"path outside the drive and this task's files: {raw}")
        if not path.exists():
            raise PermanentToolError(f"no such file: {raw}")
        return path


class ListArgs(BaseModel):
    folder: str = Field("", description="Folder relative to the drive, e.g. invoices")
    pattern: str = Field("*", description="Glob, e.g. *.pdf")


class PathArgs(BaseModel):
    path: str = Field(description="Path relative to the drive, or an absolute path returned by another tool")


class TableArgs(PathArgs):
    max_rows: int = Field(500, ge=1, le=5000)
    sheet: str | None = Field(None, description="Worksheet name (xlsx only)")


class _FileTool(Tool):
    app: ClassVar[str] = "files"
    risk: ClassVar[str] = "read"
    kind: ClassVar[str] = "file"

    def __init__(self, roots: FileRoots):
        self.roots = roots


class ListFiles(_FileTool):
    name = "files.list"
    description = "List files in a drive folder."
    input_schema = ListArgs

    async def run(self, args: ListArgs, ctx: RunContext) -> ToolResult:
        folder = self.roots.resolve(args.folder or ".", ctx)
        if not folder.is_dir():
            raise PermanentToolError(f"not a folder: {args.folder}")
        files = sorted(p for p in folder.glob(args.pattern) if p.is_file())
        return ToolResult(ok=True, output=[{"path": str(p.relative_to(self.roots.drive_root))
                                            if p.is_relative_to(self.roots.drive_root) else str(p),
                                            "size": p.stat().st_size} for p in files])


class ReadPdf(_FileTool):
    name = "files.read_pdf"
    description = "Extract text from a PDF, page by page. Untrusted content."
    input_schema = PathArgs

    async def run(self, args: PathArgs, ctx: RunContext) -> ToolResult:
        path = self.roots.resolve(args.path, ctx)
        try:
            pages = [page.extract_text() or "" for page in PdfReader(path).pages]
        except PdfReadError as exc:
            raise PermanentToolError(f"not a readable PDF: {args.path} ({exc})") from exc
        return ToolResult(ok=True, output={"path": str(path), "pages": pages}, untrusted=True)


class ReadText(_FileTool):
    name = "files.read_text"
    description = "Read a text or markdown file. Untrusted content."
    input_schema = PathArgs

    async def run(self, args: PathArgs, ctx: RunContext) -> ToolResult:
        path = self.roots.resolve(args.path, ctx)
        if path.suffix.lower() not in {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}:
            raise PermanentToolError(f"not a text file: {path.name}")
        return ToolResult(ok=True, output={"path": str(path), "text": path.read_text(errors="replace")[:20000]},
                          untrusted=True)


class ReadTable(_FileTool):
    name = "files.read_table"
    description = "Read a CSV or XLSX file as header + rows (strings). Untrusted content."
    input_schema = TableArgs

    async def run(self, args: TableArgs, ctx: RunContext) -> ToolResult:
        path = self.roots.resolve(args.path, ctx)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(path.read_text(errors="replace"))))
        elif suffix == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[args.sheet] if args.sheet else wb.active
            rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
            wb.close()
        else:
            raise PermanentToolError(f"not a CSV or XLSX file: {path.name}")
        if not rows:
            return ToolResult(ok=True, output={"path": str(path), "header": [], "rows": [], "truncated": False})
        header, body = rows[0], rows[1:]
        return ToolResult(ok=True, untrusted=True, output={
            "path": str(path), "header": header, "rows": body[:args.max_rows],
            "row_count": len(body), "truncated": len(body) > args.max_rows})


FILE_TOOLS = (ListFiles, ReadPdf, ReadText, ReadTable)


def register_file_tools(registry: ToolRegistry, roots: FileRoots) -> None:
    for cls in FILE_TOOLS:
        registry.register(cls(roots))
