"""The finance team's AR tracker workbook, as it looks before Arya starts (var/drive/finance/).

One row per open or disputed invoice. Owner and Notes are the team's own columns (some already
filled in): Arya may update only the status columns its sheet map allows. Deterministic: re-seeding
gives the same cells.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from sandbox.common.world import World

HEADER = ["Invoice", "Customer ID", "Customer", "Invoice date", "Due date", "Amount (₹)", "Balance (₹)",
          "Status", "Promised on", "Last action", "Updated", "Owner", "Notes"]

# The team's own notes: they must survive every update Arya makes.
NOTES = {"INV-2291": ("Kavya", "Ritu says payment run is on Fridays"),
         "INV-2296": ("Kavya", "Check delivery note before chasing"),
         "INV-2248": ("Legal", "With legal team. Do not contact."),
         "INV-2259": ("Kavya", "On payment plan agreed 01-Oct"),
         "INV-2240": ("Sales", "Damaged goods claim under review")}


def generate(world: World, drive: Path) -> Path:
    folder = drive / "finance"
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Receivables"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for inv in sorted(world.invoices, key=lambda i: i.number):
        cust = world.customer(inv.customer)
        owner, note = NOTES.get(inv.number, ("Kavya", None))
        ws.append([inv.number, cust.id, cust.name, inv.issue_date.isoformat(), world.due_date(inv).isoformat(),
                   float(inv.amount), float(inv.amount), "Disputed" if inv.status == "DISPUTED" else "Open",
                   None, None, None, owner, note])
    for col, width in zip("ABCDEFGHIJKLM", [11, 11, 26, 12, 12, 12, 12, 16, 12, 40, 22, 9, 40]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    wb.properties.created = wb.properties.modified = world.start.replace(tzinfo=None)
    path = folder / "ar_tracker.xlsx"
    wb.save(path)
    return path
